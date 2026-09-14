#!/usr/bin/env python3
"""Translate Balance Changes Document.xlsx using locales/ko game strings."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell

REPO_ROOT = Path(__file__).resolve().parents[1]
LOCALES_KO = REPO_ROOT / "locales" / "ko"
LOCALES_EN = REPO_ROOT / "locales" / "en"
SOURCE_XLSX = Path(r"c:\Users\yjlw4\Downloads\Balance Changes Document.xlsx")
OUTPUT_XLSX = REPO_ROOT / "Balance Changes Document (ko).xlsx"

HEADER_MAP = {
    "Dex #": "도감번호",
    "Gen": "세대",
    "Icon": "아이콘",
    "Pokemon": "포켓몬",
    "Starter": "스타터",
    "Live Passive": "현재 패시브",
    "Replacement Passive": "변경 패시브",
    "Live Passive Ability": "현재 패시브 특성",
    "Projected Passive Ability": "예정 패시브 특성",
    "Previous Passive Ability": "이전 패시브 특성",
    "Live": "현재",
    "Changes": "변경",
    "Replacement": "변경",
    "Approval": "승인",
    "Notes": "메모",
    "Live Egg Moves": "현재 알 기술",
    "Replacement Egg Moves": "변경 알 기술",
    "Projected Egg Moves": "예정 알 기술",
    "Previous Egg Moves": "이전 알 기술",
    "Common": "커먼",
    "Rare": "레어",
    "Common A": "커먼 A",
    "Common B": "커먼 B",
    "Common C": "커먼 C",
    "Cost": "코스트",
    "Starter Cost": "스타터 코스트",
    "Projected Cost": "예정 코스트",
    "Previous Cost": "이전 코스트",
    "Live Cost": "현재 코스트",
    "Egg Tier": "알 등급",
    "Projected Egg Tier": "예정 알 등급",
    "Previous Egg Tier": "이전 알 등급",
    "Live Egg Tier": "현재 알 등급",
    "Subject to Change": "변경 가능",
    "Prospective": "예정",
    "Approved by Head": "최종 승인",
    "Red = Replacing": "빨강 = 교체",
    "Yellow = May Change": "노랑 = 변경 가능",
    "Green = No Change": "초록 = 변경 없음",
    "Color Key ->": "색상 키 ->",
    "Species": "종족",
    "GitHub Species": "GitHub 종족",
    "GitHub.Species": "GitHub.종족",
    "Passives": "패시브",
    "Passives Output": "패시브 출력",
    "Egg Moves": "알 기술",
    "Egg Moves Output": "알 기술 출력",
    "Egg Tier Output": "알 등급 출력",
    "Starter Cost Output": "스타터 코스트 출력",
    "Gen Space Inserts": "세대 공백 삽입",
    "Base Pokemon": "기본 포켓몬",
    "Pokemon Form": "포켓몬 폼",
    "Sprite": "스프라이트",
    "Sublegend?": "준전설?",
    "Legendary?": "전설?",
    "Mythical?": "환상?",
    "Species Desc": "종족 설명",
    "Type 1": "타입 1",
    "Type 2": "타입 2",
    "height": "키",
    "weight": "몸무게",
    "Ability1": "특성1",
    "Ability2": "특성2",
    "HAbility": "숨겨진 특성",
    "BST": "종족값 합",
    "HP": "HP",
    "Atk": "공격",
    "Def": "방어",
    "SpAtk": "특공",
    "SpDef": "특방",
    "Spd": "스피드",
    "Catchrate": "포획률",
    "Friendship": "친밀도",
    "Base Exp": "기본 경험치",
    "GrowthType": "성장 타입",
    "Male Ratio": "수컷 비율",
    "Genderdiffs?": "성별 차이?",
    "Regional Form?": "리전 폼?",
    "Forms?": "폼?",
    "Valid Forms": "유효 폼",
    "FormKey": "폼 키",
    "Locations": "서식지",
    "INSTRUCTIONS (02/15/26)": "안내 (02/15/26)",
    "N/A": "해당 없음",
}

EGG_TIER_MAP = {
    "Common": "커먼",
    "Rare": "레어",
    "Epic": "에픽",
    "Legendary": "레전더리",
    "Manaphy": "마나피",
}

CODE_PATTERN = re.compile(
    r"(SpeciesId\.|AbilityId\.|MoveId\.|EggTier\.|\[SpeciesId\.|: \{|\],|: \d)"
)
FORMULA_STRING_PATTERN = re.compile(r'"((?:[^"]|"")*)"')


def load_json(path: Path) -> dict:
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def build_name_map(en_path: Path, ko_path: Path, nested_key: str | None = "name") -> dict[str, str]:
    en_data = load_json(en_path)
    ko_data = load_json(ko_path)
    mapping: dict[str, str] = {}

    for key, en_val in en_data.items():
        ko_val = ko_data.get(key)
        if ko_val is None:
            continue
        if nested_key and isinstance(en_val, dict) and isinstance(ko_val, dict):
            en_name = en_val.get(nested_key)
            ko_name = ko_val.get(nested_key)
            if en_name and ko_name:
                mapping[en_name] = ko_name
        elif isinstance(en_val, str) and isinstance(ko_val, str):
            mapping[en_val] = ko_val

    return mapping


def build_form_templates() -> list[tuple[re.Pattern[str], str]]:
    en_form_battle = load_json(LOCALES_EN / "pokemon-form-battle.json")
    ko_form_battle = load_json(LOCALES_KO / "pokemon-form-battle.json")
    en_append = load_json(LOCALES_EN / "pokemon-form.json")["appendForm"]
    ko_append = load_json(LOCALES_KO / "pokemon-form.json")["appendForm"]

    patterns: list[tuple[re.Pattern[str], str]] = []

    mega_suffixes = [
        ("Mega X ", "megaX"),
        ("Mega Y ", "megaY"),
        ("Mega Z ", "megaZ"),
        ("Mega ", "mega"),
        ("G-Max Single Strike ", "gigantamaxSingle"),
        ("G-Max Rapid Strike ", "gigantamaxRapid"),
        ("G-Max ", "gigantamax"),
        ("Primal ", "primal"),
        ("Eternamax ", "eternamax"),
        ("Eternal Flower ", "eternal"),
        ("Bloodmoon ", "bloodmoon"),
        ("Alolan ", "alola"),
        ("Galarian ", "galar"),
        ("Hisuian ", "hisui"),
        ("Paldean ", "paldea"),
    ]

    for prefix, key in mega_suffixes:
        if key in en_form_battle:
            en_tpl = en_form_battle[key]
            ko_tpl = ko_form_battle[key]
            patterns.append((re.compile(rf"^{re.escape(prefix)}(.+)$"), ko_tpl))
        elif key in en_append:
            en_tpl = en_append[key]
            ko_tpl = ko_append[key]
            patterns.append((re.compile(rf"^{re.escape(prefix)}(.+)$"), ko_tpl))

    return patterns


def translate_pokemon(name: str, pokemon_map: dict[str, str], form_patterns: list[tuple[re.Pattern[str], str]]) -> str:
    if not name or not isinstance(name, str):
        return name

    stripped = name.strip()
    if stripped in pokemon_map:
        return pokemon_map[stripped]

    for pattern, ko_tpl in form_patterns:
        match = pattern.match(stripped)
        if not match:
            continue
        base_name = match.group(1).strip()
        ko_base = pokemon_map.get(base_name, base_name)
        return ko_tpl.replace("{{pokemonName}}", ko_base)

    # Handle "Name (Form)" generic pattern
    paren_match = re.match(r"^(.+?) \((.+)\)$", stripped)
    if paren_match:
        base_name, form_name = paren_match.groups()
        ko_base = pokemon_map.get(base_name.strip(), base_name.strip())
        return f"{ko_base} ({form_name})"

    return stripped


def translate_term(value: str, pokemon_map: dict[str, str], ability_map: dict[str, str], move_map: dict[str, str], form_patterns) -> str:
    if value is None:
        return value
    if not isinstance(value, str):
        return value

    text = value.strip()
    if not text:
        return value

    if text in HEADER_MAP:
        return HEADER_MAP[text]
    if text in EGG_TIER_MAP:
        return EGG_TIER_MAP[text]
    if text in pokemon_map:
        return pokemon_map[text]
    if text in ability_map:
        return ability_map[text]
    if text in move_map:
        return move_map[text]

    # Pokemon with form prefixes
    pokemon_translated = translate_pokemon(text, pokemon_map, form_patterns)
    if pokemon_translated != text:
        return pokemon_translated

    return value


def translate_formula_text(
    formula: str,
    pokemon_map: dict[str, str],
    ability_map: dict[str, str],
    move_map: dict[str, str],
    form_patterns,
    sheet_name: str,
) -> str:
    def replace_string(match: re.Match[str]) -> str:
        raw = match.group(1).replace('""', '"')
        if sheet_name == "GitHub Output" and CODE_PATTERN.search(raw):
            return match.group(0)

        translated = translate_term(raw, pokemon_map, ability_map, move_map, form_patterns)
        if translated == raw:
            translated = translate_free_text(raw, pokemon_map, ability_map, move_map, form_patterns)
        if translated == raw:
            return match.group(0)

        escaped = translated.replace('"', '""')
        return f'"{escaped}"'

    return FORMULA_STRING_PATTERN.sub(replace_string, formula)


def get_formula_text(cell: Cell) -> str | None:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return value
    if hasattr(value, "text"):
        text = getattr(value, "text", None)
        if isinstance(text, str) and text.startswith("="):
            return text
    return None


def set_formula_text(cell: Cell, formula: str) -> None:
    value = cell.value
    if hasattr(value, "text"):
        value.text = formula
    else:
        cell.value = formula


def translate_free_text(text: str, pokemon_map: dict[str, str], ability_map: dict[str, str], move_map: dict[str, str], form_patterns) -> str:
    if not text or not isinstance(text, str):
        return text

    combined = {**move_map, **ability_map}
    # Longest names first to avoid partial replacements
    for en_name in sorted(combined, key=len, reverse=True):
        if en_name in text:
            text = text.replace(en_name, combined[en_name])

    for en_name in sorted(pokemon_map, key=len, reverse=True):
        if en_name in text:
            ko_name = translate_pokemon(en_name, pokemon_map, form_patterns)
            text = text.replace(en_name, ko_name)

    return text


def should_skip_translation(cell: Cell, sheet_name: str) -> bool:
    if cell.value is None:
        return True
    if isinstance(cell.value, (int, float, bool)):
        return True
    if sheet_name == "GitHub Output":
        val = str(cell.value)
        if CODE_PATTERN.search(val):
            return True
    return False


def is_notes_column(header: str | None) -> bool:
    return header in {"Notes", "메모", "Species Desc", "종족 설명"}


def is_likely_game_term_column(header: str | None) -> bool:
    if not header:
        return False
    keywords = [
        "Pokemon", "포켓몬", "Starter", "스타터", "Passive", "패시브",
        "Egg Move", "알 기술", "Ability", "특성", "Common", "커먼", "Rare", "레어",
        "Egg Tier", "알 등급", "Species", "종족", "Ability1", "Ability2", "HAbility",
        "Type", "타입", "GrowthType", "Live", "현재", "Replacement", "변경",
        "Previous", "이전", "Projected", "예정", "Changes", "Subject to Change",
        "Prospective", "Approved by Head", "최종 승인",
    ]
    return any(k in header for k in keywords)


def looks_like_header(value: str) -> bool:
    text = value.strip()
    if not text:
        return False
    if text in HEADER_MAP or text in EGG_TIER_MAP:
        return True
    if len(text) > 60:
        return False
    keywords = [
        "Dex", "Gen", "Pokemon", "Passive", "Egg", "Cost", "Approval", "Notes",
        "Common", "Rare", "Starter", "Species", "Ability", "Type", "Live",
        "Replacement", "Previous", "Projected", "Changes", "Subject", "Prospective",
        "Approved", "Red =", "Yellow =", "Green =", "Color Key",
    ]
    return any(keyword in text for keyword in keywords)


def get_header_map(ws) -> dict[int, str | None]:
    headers: dict[int, str | None] = {}
    for row in ws.iter_rows(min_row=1, max_row=4):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            text = cell.value.strip()
            if not text or not looks_like_header(text):
                continue
            if cell.column not in headers:
                headers[cell.column] = text
    return headers


def translate_workbook() -> None:
    pokemon_map = build_name_map(LOCALES_EN / "pokemon.json", LOCALES_KO / "pokemon.json", nested_key=None)
    ability_map = build_name_map(LOCALES_EN / "ability.json", LOCALES_KO / "ability.json")
    move_map = build_name_map(LOCALES_EN / "move.json", LOCALES_KO / "move.json")
    form_patterns = build_form_templates()

    shutil.copy2(SOURCE_XLSX, OUTPUT_XLSX)
    wb = openpyxl.load_workbook(OUTPUT_XLSX)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        col_headers = get_header_map(ws)

        translatable_sheets = {
            "Balance Changes (WIP)",
            "Passive Changes (WIP)",
            "Egg Move Changes (WIP)",
            "CostTier Changes (WIP)",
            "Last Patch Changes",
            "Last Patch Passive Changes",
            "Last Patch Egg Move Changes",
            "Live Starter Data",
            "Live All Forms Data",
            "Split Passive changes 1.8.0",
            "Passive History (Split)",
            "Passive History (Starters)",
            "Pokemon",
            "EvoLine Data",
            "Move Compendium",
            "EM History",
            "CostTier History",
            "Name Extractions",
        }

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or isinstance(cell.value, (int, float, bool)):
                    continue

                formula_text = get_formula_text(cell)
                if formula_text:
                    if sheet_name == "GitHub Output" and CODE_PATTERN.search(formula_text):
                        continue
                    translated_formula = translate_formula_text(
                        formula_text,
                        pokemon_map,
                        ability_map,
                        move_map,
                        form_patterns,
                        sheet_name,
                    )
                    if translated_formula != formula_text:
                        set_formula_text(cell, translated_formula)
                    continue

                if should_skip_translation(cell, sheet_name):
                    continue

                header = col_headers.get(cell.column)
                original = cell.value
                if not isinstance(original, str):
                    continue

                if original in HEADER_MAP:
                    cell.value = HEADER_MAP[original]
                    continue

                if header and ("Egg Tier" in header or "알 등급" in header) and original in EGG_TIER_MAP:
                    cell.value = EGG_TIER_MAP[original]
                    continue

                if is_notes_column(header):
                    cell.value = translate_free_text(original, pokemon_map, ability_map, move_map, form_patterns)
                    continue

                if is_likely_game_term_column(header) or sheet_name in translatable_sheets:
                    cell.value = translate_term(original, pokemon_map, ability_map, move_map, form_patterns)

    wb.save(OUTPUT_XLSX)
    print(f"Saved: {OUTPUT_XLSX}")


if __name__ == "__main__":
    translate_workbook()
